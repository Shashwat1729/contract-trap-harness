"""Structural validity of the generated Word redline (report.generate_redlined_docx)."""
from __future__ import annotations

import re
import zipfile

from docx import Document

from src.harness.report import generate_redlined_docx

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _findings(n: int) -> list[dict]:
    return [{
        "clause_type": "Renewal Term", "risk": "High", "rule_id": "P-01", "trap_id": "P-01",
        "page": 1, "line": i + 1, "span_text": f"This Agreement renews for {24 + i} months.",
        "proposed_change": "Initial Term 12 months, renewal Term 12 months.", "rationale": "Pilot must not lock in.",
        "verification": "PASS", "reasons": [],
    } for i in range(n)]


def test_redline_docx_is_structurally_valid(tmp_path):
    out = generate_redlined_docx(None, _findings(8), output_path=tmp_path / "r.docx")
    assert out == tmp_path / "r.docx" and out.exists()
    xml = zipfile.ZipFile(out).read("word/document.xml").decode("utf-8")

    # Revision ids unique across ALL w:ins/w:del (previously millisecond-clock based).
    ids = re.findall(r"<w:(?:ins|del) [^>]*w:id=\"(\d+)\"", xml)
    assert len(ids) == 24 and len(set(ids)) == 24  # del + ins + comment-block ins per finding

    # Deleted text must be w:delText; no w:t inside a w:del.
    for block in re.findall(r"<w:del .*?</w:del>", xml, flags=re.S):
        assert "<w:delText" in block and "<w:t " not in block and "<w:t>" not in block

    # No dangling comment anchors (there is no comments.xml part).
    assert "commentReference" not in xml and "commentRangeStart" not in xml

    # ST_DateTime without fractional seconds / offsets, as Word writes it.
    assert all(re.fullmatch(r"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ", d) for d in re.findall(r'w:date="([^"]+)"', xml))

    # CT_RPr child order: color before sz before u.
    for rpr in re.findall(r"<w:ins .*?<w:rPr>(.*?)</w:rPr>", xml, flags=re.S):
        tags = re.findall(r"<w:(\w+)", rpr)
        assert tags == ["color", "sz", "u"]

    # And python-docx can read it back.
    doc = Document(str(out))
    assert any("Redlined Report" in p.text for p in doc.paragraphs)


def test_redline_docx_with_no_findings(tmp_path):
    out = generate_redlined_docx(None, [], output_path=tmp_path / "empty.docx", contract_text="Hello")
    assert out.exists()
    Document(str(out))
